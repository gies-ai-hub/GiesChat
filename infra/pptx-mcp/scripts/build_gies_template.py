"""Regenerate templates/gies.pptx as an Illinois-branded 16:9 deck template.

The shipped template is the default every deck starts from (the MCP prompt calls
create_presentation_from_template with /app/templates/gies.pptx). This script is
the source of that binary - edit the constants below and re-run it rather than
hand-editing the .pptx.

    cd infra/pptx-mcp && ./.venv/bin/python scripts/build_gies_template.py
"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

# Illini tokens, matching client/src/style.css
BLUE = "13294B"
ORANGE = "FF5F05"
ORANGE_DARK = "D94F04"
SLATE = "5B6E8C"
MIST = "9FB1C7"
INK = "1D2733"
WHITE = "FFFFFF"
TYPEFACE = "Arial"

WIDE = Emu(12192000)
TALL = Emu(6858000)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "gies.pptx"
"""Output path. The template is generated from python-pptx's stock master, so re-runs are idempotent."""

# Layouts that get the light treatment: blue title over white, orange rule beneath.
LIGHT_LAYOUTS = {
    "Title and Content",
    "Two Content",
    "Comparison",
    "Title Only",
    "Content with Caption",
    "Picture with Caption",
    "Title and Vertical Text",
    "Vertical Title and Text",
}


def _q(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{A if prefix == 'a' else P}}}{local}"


def rect_xml(shape_id: int, name: str, x: int, y: int, cx: int, cy: int, color: str) -> str:
    """A borderless filled rectangle, as raw DrawingML."""
    return (
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}">'
        f"<p:nvSpPr>"
        f'<p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr/>'
        f'<p:nvPr userDrawn="1"/>'
        f"</p:nvSpPr>"
        f"<p:spPr>"
        f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f"<a:ln><a:noFill/></a:ln>"
        f"</p:spPr>"
        f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>'
        f"</p:sp>"
    )


def prepend_shape(container, xml: str) -> None:
    """Insert a shape behind everything else already in the tree."""
    sp_tree = container.element.find(_q("p:cSld") + "/" + _q("p:spTree"))
    sp_tree.insert(2, etree.fromstring(xml))


def widen(container, scale: float) -> None:
    """Scale horizontal geometry when moving 4:3 -> 16:9. Vertical is unchanged."""
    for off in container.element.iter(_q("a:off")):
        off.set("x", str(int(int(off.get("x")) * scale)))
    for ext in container.element.iter(_q("a:ext")):
        ext.set("cx", str(int(int(ext.get("cx")) * scale)))


def set_theme(master) -> None:
    theme = next(
        rel.target_part for rel in master.part.rels.values() if "theme" in rel.reltype
    )
    root = etree.fromstring(theme.blob)

    scheme = root.find(".//" + _q("a:clrScheme"))
    palette = {
        "dk1": INK,
        "lt1": WHITE,
        "dk2": BLUE,
        "lt2": "F2F4F7",
        "accent1": ORANGE,
        "accent2": BLUE,
        "accent3": ORANGE_DARK,
        "accent4": SLATE,
        "accent5": MIST,
        "accent6": "A5A5A5",
        "hlink": ORANGE,
        "folHlink": SLATE,
    }
    scheme.set("name", "Illinois")
    for slot, value in palette.items():
        node = scheme.find(_q(f"a:{slot}"))
        node.clear()
        etree.SubElement(node, _q("a:srgbClr")).set("val", value)

    fonts = root.find(".//" + _q("a:fontScheme"))
    fonts.set("name", "Illinois")
    for group in ("a:majorFont", "a:minorFont"):
        fonts.find(_q(group)).find(_q("a:latin")).set("typeface", TYPEFACE)

    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def style_text(master) -> None:
    """Master-level title and body defaults, inherited by every layout."""
    styles = master.element.find(_q("p:txStyles"))

    title = styles.find(_q("p:titleStyle") + "/" + _q("a:lvl1pPr"))
    title.set("algn", "l")
    rpr = title.find(_q("a:defRPr"))
    rpr.set("sz", "3600")
    rpr.set("b", "1")
    fill = rpr.find(_q("a:solidFill"))
    fill.clear()
    etree.SubElement(fill, _q("a:srgbClr")).set("val", BLUE)

    body = styles.find(_q("p:bodyStyle"))
    for level, size in enumerate(("2000", "1800", "1600", "1400", "1400"), start=1):
        para = body.find(_q(f"a:lvl{level}pPr"))
        if para is None:
            continue
        colour = para.find(_q("a:buClr"))
        if colour is None:
            colour = etree.Element(_q("a:buClr"))
            para.insert(0, colour)
        colour.clear()
        etree.SubElement(colour, _q("a:srgbClr")).set("val", ORANGE)
        rpr = para.find(_q("a:defRPr"))
        rpr.set("sz", size)
        fill = rpr.find(_q("a:solidFill"))
        if fill is not None:
            fill.clear()
            etree.SubElement(fill, _q("a:srgbClr")).set("val", INK)


def recolor_placeholders(layout, color: str) -> None:
    """Override inherited text colour and left-align a layout's own placeholders."""
    for shape in layout.placeholders:
        body = shape.element.find(_q("p:txBody"))
        if body is None:
            continue
        styles = body.find(_q("a:lstStyle"))
        if styles is None:
            styles = etree.Element(_q("a:lstStyle"))
            body.insert(1, styles)
        para = etree.SubElement(styles, _q("a:lvl1pPr"))
        para.set("algn", "l")
        rpr = etree.SubElement(para, _q("a:defRPr"))
        etree.SubElement(etree.SubElement(rpr, _q("a:solidFill")), _q("a:srgbClr")).set(
            "val", color
        )


def place(layout, idx: int, x: int, y: int, cx: int, cy: int) -> None:
    for shape in layout.placeholders:
        if shape.placeholder_format.idx == idx:
            shape.left, shape.top, shape.width, shape.height = Emu(x), Emu(y), Emu(cx), Emu(cy)
            return


def main() -> None:
    prs = Presentation()
    scale = WIDE / prs.slide_width
    master = prs.slide_master

    widen(master, scale)
    for layout in master.slide_layouts:
        widen(layout, scale)
    prs.slide_width, prs.slide_height = WIDE, TALL

    set_theme(master)
    style_text(master)

    rule_y = 1_430_000
    prepend_shape(master, rect_xml(900, "Accent rule", 914400, rule_y, 2_200_000, 60_000, ORANGE))

    for layout in master.slide_layouts:
        name = layout.name
        if name == "Title Slide":
            prepend_shape(
                layout, rect_xml(902, "Accent rule", 914400, 3_700_000, 2_600_000, 74_000, ORANGE)
            )
            prepend_shape(layout, rect_xml(901, "Field", 0, 0, int(WIDE), int(TALL), BLUE))
            recolor_placeholders(layout, WHITE)
            # Scaling from 4:3 strands the stock title block; re-lay it as one left-aligned column.
            place(layout, 0, 914400, 2_300_000, 9_000_000, 1_300_000)
            place(layout, 1, 914400, 3_950_000, 9_000_000, 900_000)
        elif name == "Section Header":
            prepend_shape(layout, rect_xml(903, "Field", 0, 0, int(WIDE), int(TALL), ORANGE))
            recolor_placeholders(layout, WHITE)
        elif name in LIGHT_LAYOUTS:
            prepend_shape(
                layout, rect_xml(904, "Accent rule", 914400, rule_y, 2_200_000, 60_000, ORANGE)
            )

    prs.save(str(TEMPLATE))
    print(f"wrote {TEMPLATE} ({TEMPLATE.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
