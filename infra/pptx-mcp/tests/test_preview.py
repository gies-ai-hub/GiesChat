from pptx import Presentation
from pptx.util import Inches

import gies_preview as gp


def _deck(titles):
    """A deck using the built-in 'Title and Content' layout, one slide per title."""
    pres = Presentation()
    for title in titles:
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = f"body of {title}"
    return pres


def _png():
    """Smallest valid 1x1 PNG, as a file-like object python-pptx accepts."""
    import io, base64
    data = base64.b64decode(
        b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    return io.BytesIO(data)


def test_one_card_per_slide_numbered_from_one():
    html = gp.render_deck(_deck(["Alpha", "Beta", "Gamma"]), False)
    assert html.count('class="slide"') == 3
    assert "1 / 3" in html and "2 / 3" in html and "3 / 3" in html
    assert "0 / 3" not in html


def test_slide_text_is_rendered_and_escaped():
    html = gp.render_deck(_deck(["<script>alert(1)</script>"]), False)
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_shape_without_text_renders_a_placeholder_box():
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[6])   # blank
    slide.shapes.add_picture(_png(), Inches(1), Inches(1), Inches(2), Inches(2))
    html = gp.render_deck(pres, False)
    assert 'class="ph"' in html
    assert "picture" in html


def test_positions_come_from_the_real_geometry():
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(5), Inches(1))
    box.text_frame.text = "top left"
    html = gp.render_deck(pres, False)
    assert "left:0.00%" in html
    assert "top:0.00%" in html


def test_design_note_only_when_a_template_was_uploaded():
    assert "styled with your template" not in gp.render_deck(_deck(["A"]), False)
    assert "styled with your template" in gp.render_deck(_deck(["A"]), True)
