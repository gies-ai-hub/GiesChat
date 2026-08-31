"""The default template every deck starts from must stay 16:9 and Illini-branded.

Regenerate it with scripts/build_gies_template.py; these assertions are what a
bad regeneration trips on.
"""

from pathlib import Path

import pytest
from pptx import Presentation

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "gies.pptx"

BLUE = "13294B"
ORANGE = "FF5F05"

# The MCP prompt and design_layouts both address layouts by name.
EXPECTED_LAYOUTS = {
    "Title Slide",
    "Title and Content",
    "Section Header",
    "Two Content",
    "Comparison",
    "Title Only",
    "Blank",
    "Content with Caption",
    "Picture with Caption",
    "Title and Vertical Text",
    "Vertical Title and Text",
}


@pytest.fixture(scope="module")
def template():
    return Presentation(str(TEMPLATE))


@pytest.fixture(scope="module")
def theme_xml(template):
    master = template.slide_master
    part = next(
        rel.target_part for rel in master.part.rels.values() if "theme" in rel.reltype
    )
    return part.blob.decode()


def test_is_widescreen(template):
    assert (template.slide_width, template.slide_height) == (12192000, 6858000)


def test_starts_empty(template):
    assert len(template.slides) == 0


def test_layout_names_are_unchanged(template):
    assert {l.name for l in template.slide_master.slide_layouts} == EXPECTED_LAYOUTS


def test_theme_carries_illini_palette(theme_xml):
    assert f'<a:accent1><a:srgbClr val="{ORANGE}"/></a:accent1>' in theme_xml
    assert f'<a:dk2><a:srgbClr val="{BLUE}"/></a:dk2>' in theme_xml


def test_title_slide_is_a_blue_field(template):
    layout = next(
        l for l in template.slide_master.slide_layouts if l.name == "Title Slide"
    )
    field = next(s for s in layout.shapes if s.name == "Field")
    assert f'<a:srgbClr val="{BLUE}"/>' in field.element.xml
    assert (field.width, field.height) == (template.slide_width, template.slide_height)


def test_section_header_is_an_orange_field(template):
    layout = next(
        l for l in template.slide_master.slide_layouts if l.name == "Section Header"
    )
    field = next(s for s in layout.shapes if s.name == "Field")
    assert f'<a:srgbClr val="{ORANGE}"/>' in field.element.xml


def test_content_layouts_carry_the_accent_rule(template):
    layout = next(
        l for l in template.slide_master.slide_layouts if l.name == "Title and Content"
    )
    rule = next(s for s in layout.shapes if s.name == "Accent rule")
    assert f'<a:srgbClr val="{ORANGE}"/>' in rule.element.xml


def test_branding_sits_behind_the_placeholders(template):
    """Decoration is inserted first in the shape tree, so text is never covered."""
    layout = next(
        l for l in template.slide_master.slide_layouts if l.name == "Title Slide"
    )
    names = [s.name for s in layout.shapes]
    assert names.index("Field") < names.index("Title 1")
    assert names.index("Accent rule") < names.index("Title 1")


def test_a_built_deck_inherits_the_branding(template):
    """The path the MCP actually takes: open the template, add a titled slide."""
    prs = Presentation(str(TEMPLATE))
    layout = next(l for l in prs.slide_master.slide_layouts if l.name == "Title and Content")
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Adoption"
    assert slide.shapes.title.text == "Adoption"
    assert any(s.name == "Accent rule" for s in slide.slide_layout.shapes)
