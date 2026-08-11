import importlib
import io

import pytest
from pptx import Presentation

import gies_edit as ge


def _module(name):
    """Resolve through sys.modules — tests/test_e2e.py re-imports the stack."""
    return importlib.import_module(name)


class _App:
    def __init__(self):
        self.tools = {}

    def tool(self, *a, **k):
        def wrap(fn):
            self.tools[fn.__name__] = fn
            return fn
        return wrap


@pytest.fixture(autouse=True)
def _reset():
    state = _module("gies_state")
    state._store.clear()
    state._current.clear()
    _module("gies_auth")._user.set("alice")
    yield


def _deck(titles):
    pres = Presentation()
    for title in titles:
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text = title
    return pres


def _tools(pres=None):
    state = _module("gies_state")
    if pres is not None:
        state.ScopedPresentations()["deck-1"] = pres
    app = _App()
    ge.register_edit_tools(app, state.ScopedPresentations(),
                           state.get_current_presentation_id)
    return app.tools


def _titles(pres):
    return [s.shapes.title.text for s in pres.slides]


def test_delete_removes_the_slide():
    pres = _deck(["A", "B", "C"])
    result = _tools(pres)["delete_slide"](1)
    assert result["slide_count"] == 2
    assert _titles(pres) == ["A", "C"]


def test_delete_drops_the_relationship_so_the_saved_file_is_clean():
    pres = _deck(["A", "B", "C"])
    before = len(pres.part.rels)
    _tools(pres)["delete_slide"](1)
    assert len(pres.part.rels) == before - 1

    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)
    assert len(Presentation(buf).slides) == 2


def test_move_reorders_without_changing_the_count():
    pres = _deck(["A", "B", "C"])
    result = _tools(pres)["move_slide"](2, 0)
    assert result["slide_count"] == 3
    assert _titles(pres) == ["C", "A", "B"]


def test_move_to_the_end():
    pres = _deck(["A", "B", "C"])
    _tools(pres)["move_slide"](0, 2)
    assert _titles(pres) == ["B", "C", "A"]


@pytest.mark.parametrize("index", [-1, 3])
def test_delete_rejects_an_out_of_range_index(index):
    pres = _deck(["A", "B", "C"])
    assert "error" in _tools(pres)["delete_slide"](index)
    assert len(pres.slides) == 3


@pytest.mark.parametrize("args", [(-1, 0), (0, 3), (3, 0)])
def test_move_rejects_an_out_of_range_index(args):
    pres = _deck(["A", "B", "C"])
    assert "error" in _tools(pres)["move_slide"](*args)
    assert _titles(pres) == ["A", "B", "C"]


def test_errors_when_no_deck_is_loaded():
    tools = _tools()
    assert "error" in tools["delete_slide"](0)
    assert "error" in tools["move_slide"](0, 1)
